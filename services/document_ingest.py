"""把上传的文档和图片转为可检索文本，供 RAG 使用。"""
from __future__ import annotations

from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
KNOWLEDGE_UPLOADS = PROJECT_ROOT / "knowledge" / "uploads"

DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx", ".md", ".markdown", ".txt"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}


def kind(path: str | Path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in DOCUMENT_SUFFIXES:
        return "document"
    if suffix in IMAGE_SUFFIXES:
        return "image"
    return ""


def _target(path: Path) -> Path:
    return KNOWLEDGE_UPLOADS / f"{path.name}.txt"


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-8", "big5", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
            continue

        # Scanner PDFs commonly store each page as an embedded image.  OCR it
        # only when the page has no text layer, keeping regular PDF imports fast.
        image_parts = []
        for image in page.images:
            if image.image is not None:
                ocr_text = _ocr(image.image)
                if ocr_text:
                    image_parts.append(ocr_text)
        if image_parts:
            pages.append(f"第 {number} 页\n" + "\n".join(image_parts))
    return "\n\n".join(pages).strip()


def _extract_docx(path: Path) -> str:
    from docx import Document
    document = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation
    presentation = Presentation(str(path))
    pages: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
        if parts:
            pages.append(f"第 {number} 页\n" + "\n".join(parts))
    return "\n\n".join(pages).strip()


def _ocr(source: object) -> str:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise RuntimeError("缺少图片 OCR 依赖 rapidocr-onnxruntime") from exc
    result, _elapsed = RapidOCR()(source)
    lines = [str(item[1]).strip() for item in (result or []) if len(item) > 1 and str(item[1]).strip()]
    return "\n".join(lines).strip()


def _extract_image(path: Path) -> str:
    return _ocr(str(path))


def ingest(path: str | Path) -> dict:
    """提取文本，写入 knowledge/uploads；返回可直接显示的导入信息。"""
    path = Path(path)
    source_kind = kind(path)
    if not source_kind:
        return {"ok": False, "message": "不是可检索的文档或图片"}
    try:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = _extract_pdf(path)
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".pptx":
            text = _extract_pptx(path)
        elif suffix in IMAGE_SUFFIXES:
            text = _extract_image(path)
        else:
            text = _read_text(path)
    except Exception as exc:
        return {"ok": False, "message": f"提取失败：{type(exc).__name__}: {exc}"}
    if not text:
        message = "未提取到文字（图片或扫描件请确认清晰、方向正确）"
        return {"ok": False, "message": message}
    target = _target(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")
    return {"ok": True, "message": "已提取文本并加入资料检索", "text_chars": len(text),
            "knowledge_file": str(target.relative_to(PROJECT_ROOT)).replace("\\", "/")}


def remove(path: str | Path) -> None:
    """删除原始上传文件时同步删除其可检索文本副本。"""
    try:
        _target(Path(path)).unlink(missing_ok=True)
    except OSError:
        pass


def indexed(path: str | Path) -> bool:
    """可检索文本副本是否存在，供上传列表显示导入状态。"""
    return _target(Path(path)).is_file()


def preview_text(path: str | Path, limit: int = 16000) -> dict:
    """读取已提取文本的开头，避免预览时把大型资料完整载入内存。"""
    target = _target(Path(path))
    if not target.is_file():
        return {"text": "", "text_chars": 0, "truncated": False}
    limit = max(1, min(int(limit), 50000))
    with target.open("r", encoding="utf-8", errors="replace") as handle:
        text = handle.read(limit + 1)
    truncated = len(text) > limit
    if truncated:
        text = text[:limit]
    return {"text": text, "text_chars": len(text), "truncated": truncated}
